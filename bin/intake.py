#!/usr/bin/env python3
"""intake.py — bring a document the user already has into profile/intake/.

    intake.py <file> [--dest profile/intake]   copy the file in and write a
                                               plain-text twin beside it
    intake.py --list [--dest profile/intake]   what has been brought in so far

/setup's first real step asks for everything the user has already written
down: a current resume, their own LinkedIn profile copied as text, a cover
letter or two they liked, a few ads they'd apply to today, the LinkedIn
connections export. Every one of those is a document Claude then reads, so
the interview can confirm instead of dictate. This script is the one path a
file takes from "dropped into the chat" to "readable in the repo":

* the raw file is copied into `profile/intake/` untouched, so the entry it
  seeds can name its provenance (`source: intake/resume-2024.pdf`);
* `.pdf` and `.docx` get a `.txt` twin with their text extracted, because a
  browser-only user cannot convert files and the skill should not have to
  guess at binary content; `.txt`, `.md` and `.csv` are copied as they are;
* a LinkedIn *Connections.csv* is recognised by its `First Name` header and
  written to `profile/connections.csv` with LinkedIn's preamble notes
  dropped, the only shape the referral lookup reads.

Nothing here interprets the content. An old resume is a claim the user once
made, not evidence: every number in it still gets asked "is that a number you
could point to?" before it earns a `confidence:` (CLAUDE.md red lines 1, 2).
"""
from __future__ import annotations

import argparse
import csv
import io
import os
import re
import shutil
import sys
import unicodedata

DEFAULT_DEST = os.path.join("profile", "intake")
TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".csv"}
LINKEDIN_HEADER = "First Name"


def slug(name: str) -> str:
    """`My Résumé (final) v2.PDF` → `my-resume-final-v2.pdf`; stable, ASCII."""
    stem, ext = os.path.splitext(os.path.basename(name))
    stem = unicodedata.normalize("NFKD", stem).encode("ascii", "ignore").decode()
    stem = re.sub(r"[^A-Za-z0-9]+", "-", stem).strip("-")
    return f"{stem.lower() or 'file'}{ext.lower()}"


def extract_docx(path: str) -> str:
    try:
        import docx  # python-docx; lazy so --list and text files never need it
    except ImportError as e:  # pragma: no cover - depends on the environment
        raise RuntimeError(
            "reading .docx needs python-docx (pip install python-docx); "
            "or paste the text into the chat instead"
        ) from e
    d = docx.Document(path)
    parts: list[str] = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def extract_pdf(path: str) -> str:
    try:
        from pypdf import PdfReader
    except (KeyboardInterrupt, SystemExit):
        raise
    except BaseException as e:  # pragma: no cover - depends on the environment
        # ImportError when pypdf is missing; a pyo3 panic (not an Exception)
        # when the system cryptography package it imports is broken.
        raise RuntimeError(
            "reading .pdf needs pypdf (pip install pypdf cffi); "
            "or paste the text into the chat instead"
        ) from e
    reader = PdfReader(path)
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def tidy(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip() + "\n"


def linkedin_connections(raw: str) -> str | None:
    """The export body from LinkedIn's header row on, or None if not one."""
    lines = raw.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    for i, line in enumerate(lines):
        first = next(csv.reader(io.StringIO(line)), [])
        if first and first[0].strip() == LINKEDIN_HEADER:
            body = "\n".join(lines[i:]).strip()
            return body + "\n"
    return None


def bring_in(src: str, dest: str, root: str = ".") -> tuple[str, str | None]:
    """Copy `src` into `dest`; return (raw path, text twin or None), root-relative."""
    if not os.path.isfile(src):
        raise FileNotFoundError(f"no such file: {src}")
    os.makedirs(os.path.join(root, dest), exist_ok=True)
    name = slug(src)
    ext = os.path.splitext(name)[1]

    if ext == ".csv":
        with open(src, encoding="utf-8-sig", errors="replace") as fh:
            body = linkedin_connections(fh.read())
        if body is not None:
            out = os.path.join("profile", "connections.csv")
            with open(os.path.join(root, out), "w", encoding="utf-8", newline="") as fh:
                fh.write(body)
            return out, None

    raw = os.path.join(dest, name)
    shutil.copyfile(src, os.path.join(root, raw))
    if ext in TEXT_SUFFIXES:
        return raw, raw

    if ext == ".docx":
        text = extract_docx(src)
    elif ext == ".pdf":
        text = extract_pdf(src)
    else:
        raise RuntimeError(
            f"can't read {ext or 'that'} files; paste the text into the chat, "
            "or save it as PDF, Word or plain text first"
        )
    twin = os.path.splitext(raw)[0] + ".txt"
    with open(os.path.join(root, twin), "w", encoding="utf-8") as fh:
        fh.write(tidy(text))
    return raw, twin


def listing(dest: str, root: str = ".") -> list[str]:
    base = os.path.join(root, dest)
    if not os.path.isdir(base):
        return []
    return sorted(
        os.path.join(dest, n)
        for n in os.listdir(base)
        if not n.startswith(".") and os.path.isfile(os.path.join(base, n))
    )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Bring a document into profile/intake/.")
    parser.add_argument("file", nargs="?", help="the file to bring in")
    parser.add_argument("--list", action="store_true", help="list what has been brought in")
    parser.add_argument("--dest", default=DEFAULT_DEST, help="intake folder (repo-relative)")
    parser.add_argument("--root", default=".", help="repo root (default: cwd)")
    args = parser.parse_args(argv)

    if args.list:
        files = listing(args.dest, args.root)
        if not files:
            print("nothing brought in yet")
        for f in files:
            print(f)
        return 0
    if not args.file:
        parser.error("a file to bring in, or --list")

    try:
        raw, twin = bring_in(args.file, args.dest, args.root)
    except (FileNotFoundError, RuntimeError) as e:
        print(f"intake: {e}", file=sys.stderr)
        return 1
    if twin is None:
        print(f"saved {raw}")
    elif twin == raw:
        print(f"saved {raw}")
    else:
        print(f"saved {raw}\ntext  {twin}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
